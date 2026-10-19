"""Minimal Markdown -> .docx (headings, paragraphs, bullets, numbered lists, tables, code, **bold**,
`code`, images). Used for docs/MAIN_GOAL_SPEC.docx; PDF via LibreOffice:
  python scripts/md_to_docx.py docs/MAIN_GOAL_SPEC.md docs/MAIN_GOAL_SPEC.docx
  soffice --headless --convert-to pdf --outdir docs docs/MAIN_GOAL_SPEC.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(par, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(8.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def convert(md_path: Path, out: Path):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Carlito"; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, base = 0, md_path.parent
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            buf = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i]); i += 1
            p = doc.add_paragraph(); r = p.add_run("\n".join(buf))
            r.font.name = "DejaVu Sans Mono"; r.font.size = Pt(8)
            i += 1; continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = [ln]; i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            n = max(len(c) for c in cells)
            t = doc.add_table(rows=len(cells), cols=n); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(cells):
                for ci in range(n):
                    par = t.cell(ri, ci).paragraphs[0]
                    add_inline(par, row[ci] if ci < len(row) else "", size=8.5)
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
            doc.add_paragraph(); continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 3)); add_inline(h, m.group(2))
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1; continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", ln.strip())
        if m and (base / m.group(1)).exists():
            doc.add_picture(str(base / m.group(1)), width=Cm(17)); i += 1; continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        if m:
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            if len(m.group(1)) >= 2:
                style += " 2"
            text = m.group(3); i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip() and not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                text += " " + lines[i].strip(); i += 1
            add_inline(doc.add_paragraph(style=style), text); continue
        if ln.strip() in ("---", ""):
            i += 1; continue
        text = ln.strip(); i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||```|\s*([-*]|\d+\.)\s|!\[)", lines[i]):
            text += " " + lines[i].strip(); i += 1
        add_inline(doc.add_paragraph(), text.lstrip("> "))
    doc.save(out)


if __name__ == "__main__":
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
